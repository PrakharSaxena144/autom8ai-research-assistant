"""Generate the sample knowledge base in data/corpus/.

The corpus describes a fictional company, Nimbus Robotics GmbH. It is written so that
one-shot "embed the question, take top-k, ask the LLM" retrieval fails on realistic questions:

* facts are split across documents and joined by internal codes (SUP-117, LX-9, IR-2025-031)
* a 2025 memo supersedes parts of the 2024 handbook, and board notes supersede a date in the annual report
* "Atlas" is both a robot and the internal HR portal
* one PDF page is a scanned image and the incident report exists only as a photo (OCR required)
* some answers need arithmetic over a table, and one needs outside (web) knowledge

Dev-only dependencies: reportlab, python-docx, pillow  (pip install -r requirements-dev.txt)
Run:  python scripts/generate_corpus.py
"""

from __future__ import annotations

import csv
import io
import json
import random
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT = Path(__file__).resolve().parent.parent / "data" / "corpus"
FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "/Library/Fonts", "C:/Windows/Fonts"]


def _font(name: str, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for d in FONT_DIRS:
        p = Path(d) / name
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def _scan_effect(img: Image.Image, seed: int) -> Image.Image:
    """Make a clean render look like a photocopy: tilt, grain, slight blur."""
    rnd = random.Random(seed)
    img = img.rotate(rnd.uniform(-0.8, 0.8), expand=False, fillcolor=(246, 244, 238))
    px = img.load()
    w, h = img.size
    for _ in range(w * h // 60):
        x, y = rnd.randrange(w), rnd.randrange(h)
        g = rnd.randint(170, 230)
        px[x, y] = (g, g, g)
    return img.filter(ImageFilter.GaussianBlur(0.5))


def _render_typed_page(lines: list[tuple[str, str]], size=(1654, 2339), seed=1) -> Image.Image:
    """lines: (style, text) where style in {"h1", "h2", "body", "gap"}."""
    img = Image.new("RGB", size, (246, 244, 238))
    draw = ImageDraw.Draw(img)
    fonts = {
        "h1": _font("DejaVuSans-Bold.ttf", 46),
        "h2": _font("DejaVuSans-Bold.ttf", 34),
        "body": _font("DejaVuSansMono.ttf", 30),
    }
    y = 150
    for style, text in lines:
        if style == "gap":
            y += 30
            continue
        f = fonts[style]
        draw.text((140, y), text, font=f, fill=(35, 35, 40))
        y += int(f.size * 1.55)
    return _scan_effect(img, seed)


# --------------------------------------------------------------------------------------
# 1. Annual report (PDF, last page is a scanned image with no text layer)
# --------------------------------------------------------------------------------------
def build_annual_report() -> None:
    styles = getSampleStyleSheet()
    h1, h2, body = styles["Title"], styles["Heading2"], styles["BodyText"]
    body.leading = 15
    story: list = []

    def p(text: str) -> None:
        story.append(Paragraph(text, body))
        story.append(Spacer(1, 6))

    def table(rows: list[list[str]], widths: list[float]) -> None:
        t = Table(rows, colWidths=widths)
        t.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
                    ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.grey),
                    ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 10))

    story.append(Paragraph("Nimbus Robotics GmbH - Annual Report FY2025", h1))
    p("Fiscal year 2025 covers the period 1 April 2024 to 31 March 2025. All amounts are in millions of euros "
      "(EUR m) unless stated otherwise. Nimbus Robotics GmbH is headquartered in Berlin, Germany, and was founded in 2016.")

    story.append(Paragraph("1. Letter from the CEO", h2))
    p("Dear shareholders, FY2025 was the year in which Nimbus became a two-product company in every sense. "
      "Our Atlas heavy-payload platform moved from early adopters to mainstream cold-chain and automotive customers, "
      "and our software and services business grew faster than hardware for the second year running. "
      "We ended the year with 310 active customers across 19 countries.")
    p("We also learned hard lessons. A field incident with Atlas cold-storage units in February 2025 showed that "
      "our validation at sub-zero temperatures was not good enough, and it exposed how dependent we are on single "
      "suppliers for a few critical components. Both issues are addressed in the risk section of this report. "
      "Our collaborative arm, Kestrel, completed 38 pilot installations and general availability is planned for FY2026. "
      "- Dr. Mira Okafor, Chief Executive Officer")

    story.append(Paragraph("2. Products", h2))
    p("<b>Porter (model AMR-400)</b> is our autonomous mobile robot for tote and pallet moves up to 400 kg. "
      "Launched in 2019, it remains our volume product, with 2,390 units shipped in FY2025.")
    p("<b>Atlas (model AMR-700)</b> carries loads up to 700 kg and is available in a standard version and a Cold "
      "Storage (CS) version for freezer warehouses. Atlas is certified to ISO 3691-4. We shipped 905 Atlas units "
      "in FY2025, up from 610 in FY2024.")
    p("<b>Kestrel (model CR-12)</b> is a six-axis collaborative arm with a 12 kg payload that can be mounted on Porter "
      "for mobile picking. Kestrel is in pilot deployments; it is not yet generally available.")
    p("<b>Nimbus Fleet Console</b> is our cloud software for fleet management, over-the-air firmware updates and "
      "analytics. It is sold as an annual subscription per robot.")

    story.append(Paragraph("3. Financial performance", h2))
    p("Revenue by segment and key figures for the last two fiscal years:")
    table(
        [
            ["EUR m", "FY2024", "FY2025"],
            ["Hardware revenue", "151.2", "184.9"],
            ["Software & services revenue", "31.2", "46.7"],
            ["Total revenue", "182.4", "231.6"],
            ["Gross margin", "38.1%", "40.6%"],
            ["Operating profit (EBIT)", "9.1", "19.5"],
            ["Research & development expense", "27.4", "33.8"],
        ],
        [7.5 * cm, 3 * cm, 3 * cm],
    )
    p("By region, EMEA contributed 58% of FY2025 revenue, North America 31% and Asia-Pacific 11%. "
      "Our largest customer, Frostline Logistics, accounted for 9% of FY2025 revenue.")
    p("Unit shipments: Porter 2,140 (FY2024) and 2,390 (FY2025); Atlas 610 (FY2024) and 905 (FY2025); "
      "Kestrel 0 (FY2024) and 38 pilot units (FY2025).")

    story.append(PageBreak())
    story.append(Paragraph("4. People", h2))
    p("On 31 March 2025 Nimbus employed 1,140 people: 610 at the Berlin headquarters, 290 at our engineering "
      "centre in Pune, India, and 240 in Austin, USA, which hosts sales, field service and our North American "
      "refurbishment line. Voluntary attrition was 9.8% (FY2024: 12.1%).")

    story.append(Paragraph("5. Risk factors", h2))
    p("<b>5.1 Supply concentration.</b> Several components are single-sourced. The most significant is the LX-9 3D "
      "lidar module used in every Atlas unit, which is supplied by a single supplier (internal supplier reference "
      "SUP-117). An interruption would halt Atlas production within approximately six weeks. A second source is "
      "under evaluation. The motor controller board for Kestrel (supplier reference SUP-226) is also single-sourced. "
      "Porter does not depend on any single-sourced component.")
    p("<b>5.2 Product safety and field reliability.</b> In February 2025, Atlas CS units at a customer freezer site "
      "experienced lidar dropouts that caused emergency stops. There were no injuries. The root cause and corrective "
      "firmware are documented in field incident report IR-2025-031.")
    p("<b>5.3 Customer concentration.</b> Frostline Logistics represented 9% of revenue. Loss of this customer would "
      "materially affect results.")
    p("<b>5.4 Currency.</b> 31% of revenue is earned in US dollars while most costs are in euros and Indian rupees.")

    story.append(Paragraph("6. Outlook for FY2026", h2))
    p("For FY2026 (1 April 2025 to 31 March 2026) we expect revenue growth of 18 to 22 percent, driven by Atlas and "
      "Fleet Console subscriptions, and an EBIT margin of at least 9 percent. We plan to launch Kestrel for general "
      "availability during FY2026.")
    p("Management board: Dr. Mira Okafor (CEO), Jonas Weller (CFO), Priyanka Deshmukh (CTO). "
      "Chair of the supervisory board: Helga Brandauer.")

    # Scanned auditor page (image only: must be OCR'd)
    scan = _render_typed_page(
        [
            ("h1", "Independent Auditor's Report"),
            ("gap", ""),
            ("body", "To Nimbus Robotics GmbH, Berlin"),
            ("gap", ""),
            ("body", "We have audited the financial statements of"),
            ("body", "Nimbus Robotics GmbH for the fiscal year from"),
            ("body", "1 April 2024 to 31 March 2025."),
            ("gap", ""),
            ("body", "In our opinion the financial statements give a"),
            ("body", "true and fair view of the net assets, financial"),
            ("body", "position and results of operations of the"),
            ("body", "company. Our opinion is unqualified."),
            ("gap", ""),
            ("body", "Auditor: Keller & Brandt"),
            ("body", "Wirtschaftspruefungsgesellschaft mbH, Hamburg"),
            ("body", "Signed: Dr. Tobias Keller, lead auditor"),
            ("body", "Date: 12 June 2025"),
        ],
        seed=7,
    )
    buf = io.BytesIO()
    scan.save(buf, format="PNG")
    buf.seek(0)
    story.append(PageBreak())
    story.append(RLImage(buf, width=17 * cm, height=24 * cm))

    doc = SimpleDocTemplate(str(OUT / "Nimbus_Annual_Report_FY2025.pdf"), pagesize=A4,
                            title="Nimbus Robotics Annual Report FY2025", author="Nimbus Robotics GmbH")
    doc.build(story)


# --------------------------------------------------------------------------------------
# 2. Employee handbook (DOCX)
# --------------------------------------------------------------------------------------
def build_handbook() -> None:
    d = DocxDocument()
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading("Nimbus Robotics Employee Handbook", 0)
    d.add_paragraph("Version 2024.1, effective 1 January 2024. Applies to all employees in Berlin, Pune and Austin "
                    "unless local law requires otherwise. Owner: People & Culture.")

    d.add_heading("1. Working hours", 1)
    d.add_paragraph("The standard working week is 40 hours. Core hours, when everyone is expected to be available, "
                    "are 10:00 to 15:00 local time. Overtime is compensated as time off in lieu, agreed with your manager "
                    "within three months.")

    d.add_heading("2. Remote work (policy HR-07, version 2.1)", 1)
    d.add_paragraph("All employees may work remotely for up to 2 days per week, agreed with their team. "
                    "Working from another country is allowed for a maximum of 20 working days per calendar year and "
                    "requires written approval from People & Culture because of tax and insurance rules. "
                    "New employees receive a one-time home office stipend of EUR 600 (or local equivalent).")

    d.add_heading("3. Leave", 1)
    d.add_paragraph("Annual leave is 28 working days per year. Unused leave of up to 5 days may be carried over to "
                    "31 March of the following year.")
    d.add_paragraph("Parental leave: the primary caregiver receives 16 weeks of fully paid parental leave; the "
                    "secondary caregiver receives 4 weeks of fully paid leave. Statutory parental leave beyond this "
                    "is unaffected.")
    d.add_paragraph("Sick leave: notify your manager on the first day. A doctor's certificate is required from the "
                    "fourth consecutive day of absence.")

    d.add_heading("4. Travel and expenses", 1)
    d.add_paragraph("Book economy class for flights under 6 hours. Business class is allowed for flights of 6 hours "
                    "or more with approval from a vice president. Rail is preferred for trips under 4 hours. "
                    "The daily meal allowance (per diem) is EUR 45 or local equivalent.")
    d.add_paragraph("Hotel caps per night:")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Destination", "Cap per night", "Notes"
    for row in [
        ("Berlin", "EUR 160", "Use the corporate rate at partner hotels"),
        ("Pune", "INR 9,000", "Company guest house preferred when available"),
        ("Austin", "USD 220", "Includes taxes"),
        ("Other cities", "EUR 180", "VP approval above cap"),
    ]:
        c = t.add_row().cells
        for i, v in enumerate(row):
            c[i].text = v
    d.add_paragraph("Submit expenses within 30 days through the Atlas HR portal.")

    d.add_heading("5. Learning and development", 1)
    d.add_paragraph("Every employee has a learning budget of EUR 1,500 per calendar year for courses, conferences "
                    "and certifications. Budgets do not roll over.")

    d.add_heading("6. Atlas HR portal and IT", 1)
    d.add_paragraph("Atlas is our internal HR portal (not to be confused with the Atlas robot). Use Atlas to request "
                    "leave, submit expenses, update bank details and download payslips. Access Atlas through single "
                    "sign-on at hr.nimbus.internal; multi-factor authentication is mandatory. If you are locked out, "
                    "contact the IT service desk on extension 4400.")
    d.add_paragraph("Laptops are refreshed every 3 years. Personal devices may not be used to access customer data.")

    d.add_heading("7. Test hall safety", 1)
    d.add_paragraph("Anyone entering the robot test hall in Berlin or Pune must have completed robot safety training "
                    "(course RS-101) within the last 12 months. Safety shoes and high-visibility vests are mandatory "
                    "in marked zones.")
    d.save(str(OUT / "Employee_Handbook_2024.docx"))


# --------------------------------------------------------------------------------------
# 3. Policy update memo (Markdown) - supersedes parts of the handbook
# --------------------------------------------------------------------------------------
MEMO = """# Memo: People policy update 2025

**From:** People & Culture
**Date:** 20 May 2025
**Effective:** 1 July 2025
**Supersedes:** sections 2 (Remote work) and 3 (Parental leave) of Employee Handbook version 2024.1

## Remote work: policy HR-07 version 3.0

Feedback from the 2024 engagement survey showed that one rule for everyone did not fit how teams work. From 1 July 2025:

- Engineering and product roles may work remotely up to 3 days per week.
- Operations and field service roles (test hall, refurbishment line, customer sites) may work remotely at most 1 day per week.
- All other roles keep the previous allowance of 2 days per week.

Working from another country is extended from 20 to 30 working days per calendar year, still with written approval.

## Parental leave

- Primary caregiver: 20 weeks fully paid (previously 16 weeks).
- Secondary caregiver: 6 weeks fully paid (previously 4 weeks).

Employees whose leave starts before 1 July 2025 stay on the previous entitlement unless their leave extends past that date, in which case the new entitlement applies to the full period.

## What does not change

Learning budget, travel rules and the home office stipend stay as described in the handbook.

Questions: ask People & Culture through the Atlas HR portal.
"""


# --------------------------------------------------------------------------------------
# 4. Supplier register (CSV)
# --------------------------------------------------------------------------------------
SUPPLIERS = [
    ["supplier_id", "supplier_name", "country", "component", "component_code", "used_in",
     "contract_end", "single_source", "risk_rating"],
    ["SUP-104", "Kraftwerk Zellen AG", "Germany", "LFP battery pack 105Ah", "BAT-105", "Porter", "2027-03-31", "No", "Low"],
    ["SUP-109", "Voltaris Energy Co.", "South Korea", "LFP battery pack 160Ah", "BAT-160", "Atlas", "2026-09-30", "No", "Medium"],
    ["SUP-117", "Lumenar Optics GmbH", "Austria", "3D lidar module", "LX-9", "Atlas", "2026-11-30", "Yes", "High"],
    ["SUP-121", "Precision Drives s.r.o.", "Czech Republic", "Hub motor 2kW", "HM-2K", "Porter; Atlas", "2028-01-31", "No", "Low"],
    ["SUP-133", "Harmonic Joint Systems", "Japan", "Strain wave gearbox", "SWG-14", "Kestrel", "2027-06-30", "Yes", "Medium"],
    ["SUP-208", "Brightline Sensors Ltd.", "Taiwan", "2D safety lidar", "LX-5", "Porter", "2027-12-31", "No", "Low"],
    ["SUP-214", "Castor & Wheel Nederland B.V.", "Netherlands", "Polyurethane drive wheel", "PW-200", "Porter; Atlas", "2026-05-31", "No", "Low"],
    ["SUP-219", "NordicCompute AB", "Sweden", "Edge compute module", "ECM-4", "Porter; Atlas; Kestrel", "2027-09-30", "No", "Medium"],
    ["SUP-226", "Shenzhen Anlun Electronics", "China", "Motor controller board", "MCB-3", "Kestrel", "2026-12-31", "Yes", "High"],
    ["SUP-230", "Alpen Kabel GmbH", "Germany", "Wiring harness", "WH-7", "Porter; Atlas", "2029-02-28", "No", "Low"],
]


# --------------------------------------------------------------------------------------
# 5. Product specifications (HTML)
# --------------------------------------------------------------------------------------
SPECS_HTML = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Nimbus Robotics - Product Specifications</title>
<style>body{font-family:sans-serif;max-width:900px;margin:auto} table{border-collapse:collapse} td,th{border:1px solid #999;padding:4px 8px}</style>
<script>console.log("analytics placeholder")</script></head>
<body>
<nav>Home | Products | Support | Careers</nav>
<h1>Product specifications</h1>
<p>Data sheet revision 2025-03. Values are nominal and measured at 20 &deg;C unless noted.</p>

<h2>Autonomous mobile robots</h2>
<table>
<caption>Porter and Atlas technical data</caption>
<tr><th>Specification</th><th>Porter (AMR-400)</th><th>Atlas (AMR-700)</th></tr>
<tr><td>Rated payload</td><td>400 kg</td><td>700 kg</td></tr>
<tr><td>Maximum speed</td><td>1.8 m/s</td><td>1.5 m/s</td></tr>
<tr><td>Runtime per charge</td><td>10 h</td><td>8 h</td></tr>
<tr><td>Charge time to 80%</td><td>1.5 h</td><td>2 h</td></tr>
<tr><td>Battery</td><td>48 V LFP, 105 Ah (BAT-105)</td><td>48 V LFP, 160 Ah (BAT-160)</td></tr>
<tr><td>Navigation lidar</td><td>2D safety lidar LX-5</td><td>3D lidar LX-9</td></tr>
<tr><td>Operating temperature</td><td>0 to 45 &deg;C</td><td>Standard: -5 to 40 &deg;C; Cold Storage (CS) variant: -25 to 35 &deg;C</td></tr>
<tr><td>Ingress protection</td><td>IP54</td><td>Standard IP54; CS variant IP65</td></tr>
<tr><td>Current firmware</td><td>3.9.1</td><td>4.2.3</td></tr>
<tr><td>Safety certification</td><td>CE</td><td>CE, ISO 3691-4</td></tr>
<tr><td>List price</td><td>EUR 38,500</td><td>EUR 61,000 (CS variant EUR 68,500)</td></tr>
</table>

<h3>Notes on Atlas Cold Storage</h3>
<p>The CS variant adds a heated lidar window, sealed connectors and low-temperature grease. Cold Storage units must run
firmware 4.2.3 or later when operated below -10 &deg;C. See the service bulletin for incident IR-2025-031.</p>

<h2>Collaborative arm</h2>
<table>
<tr><th>Specification</th><th>Kestrel (CR-12)</th></tr>
<tr><td>Payload</td><td>12 kg</td></tr>
<tr><td>Reach</td><td>1,300 mm</td></tr>
<tr><td>Repeatability</td><td>&plusmn;0.03 mm</td></tr>
<tr><td>Axes</td><td>6</td></tr>
<tr><td>Weight</td><td>34 kg</td></tr>
<tr><td>Availability</td><td>Pilot programme only</td></tr>
</table>

<h2>Firmware update policy</h2>
<p>Firmware is delivered over the air through Nimbus Fleet Console. Updates are staged: 5% of a fleet first, then the rest after 72 hours
without critical alerts. Customers on the Standard support plan receive security updates for 5 years after the last sale of a model.</p>
<footer>&copy; 2025 Nimbus Robotics GmbH</footer>
</body></html>
"""


# --------------------------------------------------------------------------------------
# 6. Board meeting notes (TXT)
# --------------------------------------------------------------------------------------
BOARD_NOTES = """NIMBUS ROBOTICS GMBH
SUPERVISORY BOARD MEETING - MINUTES (EXTRACT)
Date: 26 June 2025, Berlin
Present: H. Brandauer (Chair), M. Okafor (CEO), J. Weller (CFO), P. Deshmukh (CTO), 3 further board members

1. FY2025 CLOSING AND AUDIT
The CFO presented the audited FY2025 accounts. The auditor issued an unqualified opinion on 12 June 2025.
The board approved the accounts. No dividend will be paid; profits are retained for growth.

2. Q1 FY2026 TRADING UPDATE
Q1 FY2026 (April to June 2025) revenue is tracking at EUR 61.3m, slightly above plan.
Atlas order intake is strong in cold-chain logistics. Porter demand is flat in EMEA.

3. SUPPLY CHAIN: LX-9 LIDAR SECOND SOURCE
Following the risk discussion in the annual report, the board approved the qualification of a second
source for the LX-9 lidar: supplier reference SUP-241, Oriel Photonics Inc. (Canada).
Qualification budget EUR 1.2m. Target completion: Q3 FY2026. Until qualification is complete,
SUP-117 remains the only approved LX-9 supplier.
Note: SUP-241 is not yet in the supplier register, which is updated after qualification.

4. KESTREL LAUNCH
The CTO reported that Kestrel pilot customers require a certified safety-rated speed and separation
monitoring function that is not ready. The board accepted the recommendation to move Kestrel general
availability from FY2026 to Q2 FY2027. Pilot installations continue.

5. HEADCOUNT
The board approved 40 additional engineering hires in Pune during FY2026, mainly for Fleet Console and
Kestrel software. Austin field service headcount remains unchanged.

6. INCIDENT FOLLOW-UP
The CTO confirmed that firmware 4.2.3 (fix for IR-2025-031) has been rolled out to 100% of Atlas CS units
in the field as of 15 June 2025. No recurrence reported.

Next meeting: 25 September 2025.
"""


# --------------------------------------------------------------------------------------
# 7. Field incident report (scanned image, PNG)
# --------------------------------------------------------------------------------------
def build_incident_scan() -> None:
    img = _render_typed_page(
        [
            ("h1", "FIELD INCIDENT REPORT  IR-2025-031"),
            ("gap", ""),
            ("body", "Date of incident: 14 February 2025"),
            ("body", "Site: Frostline Logistics, Rotterdam (NL)"),
            ("body", "Product: Atlas CS (AMR-700), 12 units"),
            ("body", "Firmware at time of incident: 4.2.0"),
            ("gap", ""),
            ("h2", "Description"),
            ("body", "Units operating in a freezer zone at -18 C"),
            ("body", "triggered repeated emergency stops. Logs show"),
            ("body", "dropouts of the LX-9 lidar after 40 to 90"),
            ("body", "minutes in the cold zone. No injuries, no"),
            ("body", "damage to goods."),
            ("gap", ""),
            ("h2", "Root cause"),
            ("body", "Frost formed on the lidar window because"),
            ("body", "firmware 4.2.0 only switched on the window"),
            ("body", "heater below -20 C."),
            ("gap", ""),
            ("h2", "Corrective action"),
            ("body", "Firmware 4.2.3 activates the heater below"),
            ("body", "+2 C. Until 4.2.3 is installed, Atlas CS"),
            ("body", "units must not be operated below -10 C."),
            ("gap", ""),
            ("body", "Report owner: Field Quality, Austin"),
            ("body", "Status: CLOSED (4.2.3 released 3 March 2025)"),
        ],
        seed=11,
    )
    img.convert("L").save(OUT / "Field_Incident_Report_IR-2025-031_scan.png", optimize=True)


# --------------------------------------------------------------------------------------
# 8. Customer support FAQ (JSON export from a help centre)
# --------------------------------------------------------------------------------------
FAQ = {
    "exported_from": "Nimbus Help Centre",
    "exported_at": "2025-07-02",
    "articles": [
        {"id": "KB-101", "question": "What warranty comes with Nimbus robots?",
         "answer": "Porter has a 24-month warranty, Atlas has a 36-month warranty and Kestrel pilot units have a "
                   "12-month warranty. Batteries are covered for 5 years or 3,000 full charge cycles, whichever comes first."},
        {"id": "KB-102", "question": "What support plans are available?",
         "answer": "Standard support is included with every robot: business-hours support and a response within "
                   "2 business days. Priority support costs EUR 1,900 per robot per year and adds 24/7 support, "
                   "a 4-hour response time and next-business-day spare parts in the EU and USA."},
        {"id": "KB-103", "question": "How do I update robot firmware?",
         "answer": "Firmware updates are pushed from Nimbus Fleet Console. Open Fleet > Updates, choose the release "
                   "and schedule a staged rollout. Robots install updates while docked on a charger."},
        {"id": "KB-104", "question": "How much does Fleet Console cost?",
         "answer": "Fleet Console Essentials is included for the first year. After that it costs EUR 900 per robot "
                   "per year. Fleet Console Analytics, with throughput dashboards and heatmaps, costs EUR 1,400 per robot per year."},
        {"id": "KB-105", "question": "Can Porter be used in cold storage?",
         "answer": "No. Porter is rated for 0 to 45 degrees C. For freezer and chilled environments use the Atlas "
                   "Cold Storage variant."},
        {"id": "KB-106", "question": "How do I request a refurbished robot?",
         "answer": "Refurbished Porter units are available in North America from our Austin refurbishment line with "
                   "a 12-month warranty. Contact your account manager."},
    ],
}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_annual_report()
    build_handbook()
    (OUT / "Policy_Update_Memo_2025.md").write_text(MEMO, encoding="utf-8")
    with open(OUT / "supplier_register_2024-10.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(SUPPLIERS)
    (OUT / "product_specifications.html").write_text(SPECS_HTML, encoding="utf-8")
    (OUT / "board_meeting_minutes_2025-06-26.txt").write_text(BOARD_NOTES, encoding="utf-8")
    build_incident_scan()
    (OUT / "support_faq_export.json").write_text(json.dumps(FAQ, indent=2), encoding="utf-8")
    for p in sorted(OUT.iterdir()):
        print(f"{p.stat().st_size:>9,}  {p.name}")


if __name__ == "__main__":
    main()
