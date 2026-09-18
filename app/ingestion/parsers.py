"""Turn raw bytes of any supported format into structured text blocks.

There is exactly one entry point, `parse_document(data, filename, content_type)`. The format is detected
from magic bytes first, then extension, then MIME type, so callers never choose a parser themselves.
Every parser returns the same `ParsedDocument` shape: blocks of text with page and section metadata.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

MIN_PAGE_CHARS = 40  # below this a PDF page is treated as scanned and OCR'd


class UnsupportedFormat(ValueError):
    pass


class ParseError(ValueError):
    pass


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class ParsedDocument:
    file_type: str
    blocks: list[Block]
    title: str | None = None
    ocr_pages: list[int] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(b.text) for b in self.blocks)


# --------------------------------------------------------------------------------------
# Detection
# --------------------------------------------------------------------------------------
EXT_MAP = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".html": "html", ".htm": "html", ".xhtml": "html",
    ".md": "markdown", ".markdown": "markdown",
    ".txt": "text", ".text": "text", ".log": "text", ".rst": "text",
    ".csv": "csv", ".tsv": "tsv",
    ".json": "json",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".tif": "image", ".tiff": "image",
    ".bmp": "image", ".webp": "image", ".gif": "image",
}
MIME_MAP = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/html": "html", "application/xhtml+xml": "html",
    "text/markdown": "markdown",
    "text/plain": "text",
    "text/csv": "csv", "text/tab-separated-values": "tsv",
    "application/json": "json",
}


def detect_type(data: bytes, filename: str = "", content_type: str | None = None) -> str:
    head = data[:16]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG") or head.startswith(b"\xff\xd8\xff") or head[:6] in (b"GIF87a", b"GIF89a"):
        return "image"
    if head.startswith(b"II*\x00") or head.startswith(b"MM\x00*") or (head[:4] == b"RIFF" and data[8:12] == b"WEBP"):
        return "image"
    if head.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
        raise UnsupportedFormat("ZIP-based file that is not a .docx (e.g. .xlsx/.pptx) is not supported yet")

    ext = Path(filename or "").suffix.lower()
    if ext in EXT_MAP:
        return EXT_MAP[ext]
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime in MIME_MAP:
        return MIME_MAP[mime]
    if mime.startswith("image/"):
        return "image"

    # Unknown: accept it if it decodes as mostly printable text.
    text = _decode(data)
    sample = text[:2000]
    if sample and sum(ch.isprintable() or ch in "\n\r\t" for ch in sample) / len(sample) > 0.95:
        if re.search(r"<(html|body|p|div)\b", sample, re.IGNORECASE):
            return "html"
        return "text"
    raise UnsupportedFormat(f"Could not recognise the format of '{filename or 'upload'}'")


# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


def _ocr_image(img) -> str:
    import pytesseract
    from PIL import ImageOps

    img = ImageOps.exif_transpose(img).convert("L")
    if img.width < 1200:  # small images OCR badly; upscale
        scale = 1600 / img.width
        img = img.resize((int(img.width * scale), int(img.height * scale)))
    return pytesseract.image_to_string(img)


_MD_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_NUM_HEADING = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+([A-Z][^\n]{2,80})$")


def _looks_like_heading(line: str, prev_blank: bool = False, next_blank: bool = False) -> str | None:
    s = line.strip()
    if not s or len(s) > 90:
        return None
    m = _MD_HEADING.match(s)
    if m:
        return m.group(2).strip()
    if s.endswith((".", ",", ";", ":", "?", "!")) or "|" in s:
        return None
    words = s.split()
    m = _NUM_HEADING.match(s)
    if m and len(words) <= 12:
        return s
    letters = [c for c in s if c.isalpha()]
    if 2 <= len(words) <= 10 and len(letters) >= 6 and all(c.isupper() for c in letters):
        return s
    # short stand-alone Title line surrounded by blank lines (common in OCR output and plain text)
    if (prev_blank and next_blank and len(words) <= 6 and len(s) <= 50 and ":" not in s
            and s[0].isupper() and not any(ch.isdigit() for ch in s)):
        return s
    return None


def split_sections(text: str, initial_section: str | None = None) -> list[tuple[str | None, str]]:
    """Split text on heading-like lines. Returns [(section_title, body)].

    Consecutive headings are joined ("3. Risk factors > 3.1 Supply"); a carried-over initial
    section that has no body of its own is dropped instead of being joined.
    """
    lines = text.split("\n")
    sections: list[tuple[str | None, list[str], bool]] = [(initial_section, [], False)]
    for i, line in enumerate(lines):
        prev_blank = i == 0 or not lines[i - 1].strip()
        next_blank = i + 1 >= len(lines) or not lines[i + 1].strip()
        heading = _looks_like_heading(line, prev_blank, next_blank)
        if heading:
            sections.append((heading, [], True))
        else:
            sections[-1][1].append(line)

    merged: list[tuple[str | None, str]] = []
    pending: str | None = None
    for title, body_lines, is_heading in sections:
        body = clean_text("\n".join(body_lines))
        if not body:
            if is_heading:
                pending = f"{pending} > {title}" if pending else title
            continue
        if pending:
            title = f"{pending} > {title}" if title else pending
            pending = None
        merged.append((title, body))
    return merged


def rows_to_lines(rows: list[list[str]], caption: str | None = None, row_labels: bool = True) -> list[str]:
    """Linearise a table so every line keeps its column context.

    With row_labels=True (spec tables):  "Rated payload: Porter (AMR-400) = 400 kg; Atlas (AMR-700) = 700 kg"
    With row_labels=False (records):     "supplier_id = SUP-117; supplier_name = Lumenar Optics GmbH; ..."
    """
    rows = [[re.sub(r"\s+", " ", c or "").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    lines = [f"Table: {caption}"] if caption else []
    if not body:
        return lines + [" | ".join(header)]
    for r in body:
        if len(r) != len(header):
            lines.append(" | ".join(c for c in r if c))
            continue
        if row_labels and len(r) > 1:
            pairs = [f"{header[i]} = {r[i]}" for i in range(1, len(r)) if r[i]]
            lines.append(f"{r[0]}: " + "; ".join(pairs))
        else:
            lines.append("; ".join(f"{header[i]} = {r[i]}" for i in range(len(r)) if r[i]))
    return lines


# --------------------------------------------------------------------------------------
# Parsers
# --------------------------------------------------------------------------------------
_COLUMN_GAP = re.compile(r"(?<=\S) {5,}(?=\S)")


def _extract_pdf_page_text(page) -> str:
    """Layout mode keeps table rows on one line; wide gaps become ' | ' column separators."""
    try:
        text = page.extract_text(extraction_mode="layout") or ""
        return "\n".join(_COLUMN_GAP.sub(" | ", line) for line in text.split("\n"))
    except Exception:
        try:
            return page.extract_text() or ""
        except Exception:
            return ""


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Unreadable PDF: {exc}") from exc

    blocks: list[Block] = []
    ocr_pages: list[int] = []
    pdfium_doc = None
    for index, page in enumerate(reader.pages):
        page_no = index + 1
        text = _extract_pdf_page_text(page)
        if len(text.strip()) < MIN_PAGE_CHARS and ocr_available():
            try:
                import pypdfium2 as pdfium

                pdfium_doc = pdfium_doc or pdfium.PdfDocument(data)
                image = pdfium_doc[index].render(scale=2.5).to_pil()
                ocr_text = _ocr_image(image)
                if len(ocr_text.strip()) > len(text.strip()):
                    text = ocr_text
                    ocr_pages.append(page_no)
            except Exception as exc:
                log.warning("OCR failed on PDF page %s: %s", page_no, exc)
        text = clean_text(text)
        if not text:
            continue
        # sections can span pages, so carry the last heading forward (not onto scanned pages,
        # which are usually separate documents such as signed letters)
        current = blocks[-1].section if blocks and page_no not in ocr_pages else None
        for sec, body in split_sections(text, initial_section=current):
            blocks.append(Block(text=body, page=page_no, section=sec))

    title = None
    try:
        if reader.metadata and reader.metadata.title:
            title = str(reader.metadata.title)
    except Exception:
        pass
    return ParsedDocument(file_type="pdf", blocks=blocks, title=title, ocr_pages=ocr_pages)


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Unreadable DOCX: {exc}") from exc

    blocks: list[Block] = []
    title: str | None = None
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            text = clean_text("\n".join(buffer))
            if text:
                blocks.append(Block(text=text, section=section))
            buffer.clear()

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = (item.style.name if item.style is not None else "").lower()
            if style == "title":
                title = title or text
                flush()
                section = text
            elif style.startswith("heading"):
                flush()
                section = text
            else:
                buffer.append(text)
        elif isinstance(item, Table):
            rows = []
            for row in item.rows:
                cells, last = [], None
                for cell in row.cells:  # merged cells repeat; drop consecutive duplicates
                    if cell._tc is not last:
                        cells.append(cell.text)
                    last = cell._tc
                rows.append(cells)
            buffer.extend(rows_to_lines(rows))
    flush()
    core_title = (document.core_properties.title or "").strip()
    return ParsedDocument(file_type="docx", blocks=blocks, title=title or core_title or None)


def _parse_html(data: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "form", "iframe", "head"]):
        tag.decompose()

    for table in soup.find_all("table"):
        caption = table.find("caption")
        cap = caption.get_text(" ", strip=True) if caption else None
        rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in table.find_all("tr")]
        replacement = soup.new_tag("pre")
        replacement.string = "\n" + "\n".join(rows_to_lines(rows, caption=cap)) + "\n"
        table.replace_with(replacement)

    for level in range(1, 7):
        for h in soup.find_all(f"h{level}"):
            h.replace_with(f"\n{'#' * level} {h.get_text(' ', strip=True)}\n")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for tag in soup.find_all(["p", "li", "div", "section", "article", "pre", "tr", "dd", "dt", "blockquote"]):
        tag.insert_after("\n")

    root = soup.body or soup
    text = clean_text(root.get_text())
    blocks = [Block(text=body, section=sec) for sec, body in split_sections(text)]
    return ParsedDocument(file_type="html", blocks=blocks, title=title)


def _parse_text(data: bytes, file_type: str) -> ParsedDocument:
    text = clean_text(_decode(data))
    title = None
    first = text.split("\n", 1)[0] if text else ""
    m = _MD_HEADING.match(first)
    if m:
        title = m.group(2).strip()
    blocks = [Block(text=body, section=sec) for sec, body in split_sections(text)]
    return ParsedDocument(file_type=file_type, blocks=blocks, title=title)


def _parse_delimited(data: bytes, file_type: str) -> ParsedDocument:
    text = _decode(data)
    delimiter = "\t" if file_type == "tsv" else ","
    if file_type == "csv":
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            pass
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    lines = rows_to_lines(rows, row_labels=False)
    if not lines:
        raise ParseError("Empty table")
    header = ", ".join(rows[0])
    # keep rows in groups so a chunk never loses the column list
    blocks = []
    group = 25
    for start in range(0, len(lines), group):
        body = "\n".join(lines[start : start + group])
        blocks.append(Block(text=f"Columns: {header}\n{body}", section=f"rows {start + 1}-{min(start + group, len(lines))}"))
    return ParsedDocument(file_type=file_type, blocks=blocks)


def _parse_json(data: bytes) -> ParsedDocument:
    raw = _decode(data)
    try:
        obj = json.loads(raw)
    except json.JSONDecodeError:
        return _parse_text(data, "text")

    blocks: list[Block] = []

    def is_record_list(v) -> bool:
        return isinstance(v, list) and v and all(isinstance(x, dict) for x in v)

    def record_text(rec: dict) -> str:
        return "\n".join(f"{k}: {v if not isinstance(v, (dict, list)) else json.dumps(v, ensure_ascii=False)}"
                         for k, v in rec.items())

    if isinstance(obj, dict):
        scalars = {k: v for k, v in obj.items() if not isinstance(v, (dict, list))}
        if scalars:
            blocks.append(Block(text=record_text(scalars), section="metadata"))
        for key, value in obj.items():
            if is_record_list(value):
                for i, rec in enumerate(value, start=1):
                    blocks.append(Block(text=record_text(rec), section=f"{key} #{i}"))
            elif isinstance(value, (dict, list)):
                blocks.append(Block(text=json.dumps(value, indent=2, ensure_ascii=False), section=key))
    elif is_record_list(obj):
        for i, rec in enumerate(obj, start=1):
            blocks.append(Block(text=record_text(rec), section=f"record #{i}"))
    else:
        blocks.append(Block(text=json.dumps(obj, indent=2, ensure_ascii=False)))
    return ParsedDocument(file_type="json", blocks=[b for b in blocks if b.text.strip()])


def _parse_image(data: bytes) -> ParsedDocument:
    if not ocr_available():
        raise ParseError("Image uploads need OCR, but the tesseract binary is not installed on this server")
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Unreadable image: {exc}") from exc
    frames = []
    try:
        for i in range(getattr(img, "n_frames", 1)):  # multi-page TIFF
            img.seek(i)
            frames.append(_ocr_image(img.copy()))
    except EOFError:
        pass
    blocks = []
    for i, text in enumerate(frames, start=1):
        for sec, body in split_sections(clean_text(text)):
            blocks.append(Block(text=body, page=i if len(frames) > 1 else None, section=sec))
    return ParsedDocument(file_type="image", blocks=blocks, ocr_pages=list(range(1, len(frames) + 1)))


def parse_document(data: bytes, filename: str = "", content_type: str | None = None) -> ParsedDocument:
    if not data:
        raise ParseError("Empty file")
    kind = detect_type(data, filename, content_type)
    if kind == "pdf":
        parsed = _parse_pdf(data)
    elif kind == "docx":
        parsed = _parse_docx(data)
    elif kind == "html":
        parsed = _parse_html(data)
    elif kind in {"markdown", "text"}:
        parsed = _parse_text(data, kind)
    elif kind in {"csv", "tsv"}:
        parsed = _parse_delimited(data, kind)
    elif kind == "json":
        parsed = _parse_json(data)
    elif kind == "image":
        parsed = _parse_image(data)
    else:  # pragma: no cover
        raise UnsupportedFormat(kind)
    if not parsed.blocks or parsed.char_count < 20:
        raise ParseError(f"No extractable text found in '{filename or 'upload'}' (detected as {kind})")
    return parsed
